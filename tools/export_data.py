"""แปลงไฟล์ Excel ข้อมูลเกม -> src/data/*.json ที่หน้าเว็บ (React) import ใช้

    python tools/export_data.py

แหล่งข้อมูล (แก้ที่ Excel เท่านั้น ห้ามแก้ไฟล์ใน src/data/*.json ด้วยมือ)
ตอนรัน npm run dev ไฟล์ Excel ถูกเฝ้าดูไว้ บันทึก Excel แล้วหน้าเว็บอัปเดตเอง:
  DND_GameData.xlsx           ชีต Races, Classes, Abilities, Alignments, Settings
  DND_Skill_Extra_Final.xlsx  ชีต Normal Skills (สกิลปกติ LV0-LV3), Special Skills (สกิลพิเศษ LV0-LV5)

ทุกชีตใน DND_GameData.xlsx: แถว 1 = คำอธิบาย, แถว 2 = หัวคอลัมน์, แถว 3 ขึ้นไป = ข้อมูล
"""
import json
import re
import sys
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
GAME_XLSX = ROOT / "DND_GameData.xlsx"
SKILL_XLSX = ROOT / "DND_Skill_Extra_Final.xlsx"
OUT_DIR = ROOT / "src" / "data"
SOURCES = (GAME_XLSX, SKILL_XLSX)
SKILL_SHEETS = {"Normal Skills": "normal", "Special Skills": "special"}
OUTPUTS = (OUT_DIR / "gamedata.json", OUT_DIR / "skills.json")


class DataError(Exception):
    pass


def clean(value):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return re.sub(r"\s+", " ", str(value)).strip()


def to_int(value, where):
    text = clean(value)
    if not text:
        return 0
    try:
        return int(float(text))
    except ValueError:
        raise DataError(f"{where}: ต้องเป็นตัวเลข แต่ได้ '{text}'")


def split_list(value):
    text = clean(value)
    if not text or text.lower() == "none":
        return []
    return [part.strip() for part in text.split(",") if part.strip()]


def slug(name):
    return re.sub(r"[^a-z0-9]+", "-", name.lower().replace("’", "").replace("'", "")).strip("-")


def read_table(wb, sheet):
    """คืนรายการ dict ต่อแถว โดยใช้หัวคอลัมน์แถวที่ 2 เป็น key (ข้ามแถวที่ไม่มี ID)"""
    if sheet not in wb.sheetnames:
        raise DataError(f"ไม่พบชีต '{sheet}' ใน {GAME_XLSX.name}")
    rows = list(wb[sheet].iter_rows(min_row=2, values_only=True))
    headers = [clean(h) for h in rows[0]]
    out = []
    for n, row in enumerate(rows[1:], start=3):
        record = {h: v for h, v in zip(headers, row) if h}
        if clean(row[0]):
            record["_row"] = f"{sheet}!แถว {n}"
            out.append(record)
    return out


def check_unique(items, key, label):
    seen = set()
    for item in items:
        if item[key] in seen:
            raise DataError(f"{label}: ID '{item[key]}' ซ้ำกัน")
        seen.add(item[key])


def load_game():
    wb = openpyxl.load_workbook(GAME_XLSX, read_only=True, data_only=True)

    abilities = [
        {"id": clean(r["Code"]).upper(), "nameEn": clean(r["Name (EN)"]), "nameTh": clean(r["Name (TH)"])}
        for r in read_table(wb, "Abilities")
    ]
    codes = [a["id"] for a in abilities]
    check_unique(abilities, "id", "Abilities")
    for required in ("CON", "DEX"):
        if required not in codes:
            raise DataError(f"Abilities: ต้องมีค่าสถานะ {required} (ใช้คำนวณ HP/AC)")

    races = []
    for r in read_table(wb, "Races"):
        bonus = {}
        for code in codes:
            value = to_int(r.get(code), f"{r['_row']} คอลัมน์ {code}")
            if value:
                bonus[code] = value
        races.append({
            "id": clean(r["ID"]).lower(),
            "nameEn": clean(r["Name (EN)"]),
            "nameTh": clean(r["Name (TH)"]),
            "summary": clean(r["Description (TH)"]),
            "bonus": bonus,
            "speed": to_int(r["Speed (ft)"], f"{r['_row']} Speed"),
            "trait": {"en": clean(r["Trait (EN)"]), "th": clean(r["Trait (TH)"])},
        })
    check_unique(races, "id", "Races")

    classes = []
    for r in read_table(wb, "Classes"):
        primary = [clean(r.get(k)).upper() for k in ("Primary 1", "Primary 2") if clean(r.get(k))]
        for p in primary:
            if p not in codes:
                raise DataError(f"{r['_row']}: Primary '{p}' ไม่มีในชีต Abilities")
        hit_die = to_int(r["Hit Die"], f"{r['_row']} Hit Die")
        if hit_die <= 0:
            raise DataError(f"{r['_row']}: ต้องกรอก Hit Die")
        # ค่าสถานะที่อาชีพบวกให้ (คอลัมน์ STR..CHA ถ้าไม่มีคอลัมน์หรือเว้นว่าง = 0)
        bonus = {}
        for code in codes:
            value = to_int(r.get(code), f"{r['_row']} คอลัมน์ {code}")
            if value:
                bonus[code] = value
        classes.append({
            "id": clean(r["ID (EN)"]),
            "nameTh": clean(r["Name (TH)"]),
            "role": clean(r["Role (EN)"]),
            "hitDie": hit_die,
            "primary": primary,
            "bonus": bonus,
            "summary": clean(r["Description (TH)"]),
        })
    check_unique(classes, "id", "Classes")

    alignments = [
        {"id": clean(r["Code"]).upper(), "en": clean(r["Name (EN)"]), "th": clean(r["Name (TH)"])}
        for r in read_table(wb, "Alignments")
    ]

    raw = {clean(r["Key"]): r["Value"] for r in read_table(wb, "Settings")}
    standard = [to_int(x, "Settings standard_array") for x in split_list(raw.get("standard_array"))]
    if len(standard) != len(codes):
        raise DataError(f"Settings: standard_array ต้องมี {len(codes)} ค่าเท่ากับจำนวนค่าสถานะ (ตอนนี้มี {len(standard)})")
    settings = {
        "standardArray": sorted(standard, reverse=True),
        "maxRerolls": to_int(raw.get("max_rerolls"), "Settings max_rerolls"),
        "startingSkillLevel": to_int(raw.get("starting_skill_level"), "Settings starting_skill_level"),
        "startingSkillCount": to_int(raw.get("starting_skill_count"), "Settings starting_skill_count") or 1,
        "baseAc": to_int(raw.get("base_ac"), "Settings base_ac") or 10,
    }

    return {"abilities": abilities, "races": races, "classes": classes, "alignments": alignments, "settings": settings}


def load_skills():
    wb = openpyxl.load_workbook(SKILL_XLSX, read_only=True, data_only=True)
    skills = []
    for sheet, tier in SKILL_SHEETS.items():
        if sheet not in wb.sheetnames:
            raise DataError(f"ไม่พบชีต '{sheet}' ใน {SKILL_XLSX.name}")
        for row in wb[sheet].iter_rows(min_row=2, values_only=True):
            lv, name_en, name_th, effect, note, required, classes = (list(row) + [None] * 7)[:7]
            if lv is None or not clean(name_en):
                continue  # แถวว่าง หรือแถวหัวข้อเลเวลที่มีแค่ LV
            skills.append({
                "id": slug(clean(name_en)),
                "level": int(clean(lv).upper().replace("LV", "")),
                "tier": tier,
                "nameEn": clean(name_en),
                "nameTh": clean(name_th),
                "effect": clean(effect),
                "note": clean(note),
                "requires": split_list(required),
                "classes": split_list(classes) or ["All"],
            })
    check_unique(skills, "id", SKILL_XLSX.name)
    return skills


def warnings_for(game, skills):
    """ปัญหาที่ไม่ทำให้เว็บพัง แต่ควรรู้"""
    class_ids = {c["id"] for c in game["classes"]}
    skill_names = {s["nameEn"] for s in skills}
    out = []
    for s in skills:
        unknown = [c for c in s["classes"] if c != "All" and c not in class_ids]
        if unknown:
            out.append(f"สกิล {s['nameEn']}: อาชีพ {', '.join(unknown)} ไม่มีในชีต Classes")
        missing = [r for r in s["requires"] if r not in skill_names]
        if missing:
            out.append(f"สกิล {s['nameEn']}: สกิลที่ต้องมีก่อน {', '.join(missing)} ไม่มีในไฟล์สกิล")
    return out


def write_json(path, value):
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def export(verbose=True):
    game = load_game()
    skills = load_skills()
    OUT_DIR.mkdir(exist_ok=True)
    write_json(OUTPUTS[0], game)
    write_json(OUTPUTS[1], skills)
    warns = warnings_for(game, skills)
    if verbose:
        print(f"OK  {len(game['races'])} races, {len(game['classes'])} classes, "
              f"{len(game['abilities'])} abilities, {len(skills)} skills")
        for w in warns:
            print(f"!   {w}")
    return warns


def is_stale():
    try:
        newest_source = max(p.stat().st_mtime for p in SOURCES)
        oldest_output = min(p.stat().st_mtime for p in OUTPUTS)
    except FileNotFoundError:
        return True
    return newest_source > oldest_output


if __name__ == "__main__":
    try:
        export()
    except (DataError, KeyError) as err:
        print(f"ERROR  {err}")
        sys.exit(1)
